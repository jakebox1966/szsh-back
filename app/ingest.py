"""
국민연금 사업장 데이터 → DB 적재 스크립트. (CSV / Excel 모두 지원)

사용법:
    python -m app.ingest --file 202601_사업장.csv
    python -m app.ingest --file 202601_사업장.xlsx
    python -m app.ingest --file data.csv --min-subscribers 10
"""
import argparse
import csv
import re
import sys
from datetime import date
from pathlib import Path

from sqlalchemy.dialects.postgresql import insert as pg_insert

from app.database import engine, SessionLocal, Base
from app.models import Industry, NpsRaw, NpsMonthly

# ── 컬럼 매핑 ────────────────────────────────────────────────────

HEADER_MAP = {
    "자료생성년월": "year_month",
    "사업장명": "name",
    "사업자등록번호": "biz_no",
    "사업장가입상태코드 1 등록 2 탈퇴": "enrollment_status",
    "사업장가입상태코드": "enrollment_status",
    "우편번호": "postal_code",
    "사업장지번상세주소": "address_lot",
    "사업장도로명상세주소": "address_road",
    "고객법정동주소코드": "legal_dong_code",
    "고객행정동주소코드": "admin_dong_code",
    "법정동주소광역시도코드": "sido_code",
    "법정동주소광역시시군구코드": "sigungu_code",
    "법정동주소광역시시군구읍면동코드": "dong_code",
    "사업장형태구분코드 1 법인 2 개인": "workplace_type",
    "사업장형태구분코드": "workplace_type",
    "사업장업종코드": "industry_code",
    "사업장업종코드명": "industry_name",
    "적용일자": "applied_at",
    "재등록일자": "re_registered_at",
    "탈퇴일자": "withdrawn_at",
    "가입자수": "subscribers",
    "당월고지금액": "notified_amount",
    "신규취득자수": "new_acquisitions",
    "상실가입자수": "lost_subscribers",
}


def _normalize_header(h: str) -> str:
    return re.sub(r"\s+", " ", str(h).strip())


def _parse_date(val) -> date | None:
    if val is None or str(val).strip() == "":
        return None
    s = str(val).strip().replace(".", "-").replace("/", "-")
    parts = s.split("-")
    if len(parts) == 3:
        try:
            return date(int(parts[0]), int(parts[1]), int(parts[2]))
        except (ValueError, IndexError):
            return None
    if len(s) == 8 and s.isdigit():
        return date(int(s[:4]), int(s[4:6]), int(s[6:8]))
    return None


def _parse_year_month(val) -> date | None:
    if val is None:
        return None
    s = str(val).strip()

    # 'Jan.26', 'Feb.25'
    month_map = {
        "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
        "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
    }
    m = re.match(r"([A-Za-z]{3})\.?(\d{2})", s)
    if m:
        mon = month_map.get(m.group(1).lower())
        yr = 2000 + int(m.group(2))
        if mon:
            return date(yr, mon, 1)

    # '2026-01' or '2026.01'
    m = re.match(r"(\d{4})[-./](\d{1,2})", s)
    if m:
        return date(int(m.group(1)), int(m.group(2)), 1)

    # '202601'
    if len(s) == 6 and s.isdigit():
        return date(int(s[:4]), int(s[4:6]), 1)

    return None


def _safe_int(val) -> int | None:
    if val is None or str(val).strip() == "":
        return None
    try:
        return int(float(str(val).replace(",", "")))
    except (ValueError, TypeError):
        return None


def _safe_str(val) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    return s if s else None


# ── 파일 읽기 ────────────────────────────────────────────────────


def _read_csv(path: Path) -> tuple[list[str], list[list[str]]]:
    """CSV 파일 읽기 (cp949 → utf-8 순서로 시도)."""
    for enc in ["cp949", "euc-kr", "utf-8", "utf-8-sig"]:
        try:
            with open(path, "r", encoding=enc) as f:
                reader = csv.reader(f)
                headers = next(reader)
                rows = list(reader)
            print(f"  인코딩: {enc}")
            return headers, rows
        except (UnicodeDecodeError, UnicodeError):
            continue
    print("지원하는 인코딩을 찾을 수 없습니다 (cp949/euc-kr/utf-8)")
    sys.exit(1)


def _read_xlsx(path: Path) -> tuple[list[str], list[list]]:
    """Excel 파일 읽기."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows()
    headers = [cell.value for cell in next(rows_iter)]
    rows = [[cell.value for cell in row] for row in rows_iter]
    wb.close()
    return headers, rows


# ── 메인 로직 ────────────────────────────────────────────────────


def ingest(file_path: str, min_subscribers: int = 5):
    path = Path(file_path)
    if not path.exists():
        print(f"파일을 찾을 수 없습니다: {path}")
        sys.exit(1)

    print(f"파일 로딩 중: {path.name} ...")

    ext = path.suffix.lower()
    if ext == ".csv":
        raw_headers, raw_rows = _read_csv(path)
    elif ext in (".xlsx", ".xls"):
        raw_headers, raw_rows = _read_xlsx(path)
    else:
        print(f"지원하지 않는 파일 형식: {ext} (.csv 또는 .xlsx만 지원)")
        sys.exit(1)

    print(f"  전체 행: {len(raw_rows)}개")

    # 헤더 → 내부 키 매핑
    col_map = {}
    for idx, h in enumerate(raw_headers):
        if h is None:
            continue
        normalized = _normalize_header(h)
        for pattern, key in HEADER_MAP.items():
            if normalized == pattern:
                col_map[idx] = key
                break

    missing = {"year_month", "biz_no", "name", "subscribers"} - set(col_map.values())
    if missing:
        print(f"필수 컬럼을 찾을 수 없습니다: {missing}")
        print(f"발견된 헤더: {raw_headers}")
        sys.exit(1)

    # 행 파싱
    print("행 파싱 중...")
    records = []
    skipped = 0
    for row in raw_rows:
        cells = {}
        for i, key in col_map.items():
            cells[key] = row[i] if i < len(row) else None

        subs = _safe_int(cells.get("subscribers"))
        if subs is None or subs < min_subscribers:
            skipped += 1
            continue

        year_month = _parse_year_month(cells.get("year_month"))
        if year_month is None:
            skipped += 1
            continue

        records.append({
            "year_month": year_month,
            "biz_no": _safe_str(cells.get("biz_no")),
            "name": _safe_str(cells.get("name")),
            "enrollment_status": _safe_int(cells.get("enrollment_status")),
            "postal_code": _safe_str(cells.get("postal_code")),
            "address_lot": _safe_str(cells.get("address_lot")),
            "address_road": _safe_str(cells.get("address_road")),
            "legal_dong_code": _safe_str(cells.get("legal_dong_code")),
            "admin_dong_code": _safe_str(cells.get("admin_dong_code")),
            "sido_code": _safe_str(cells.get("sido_code")),
            "sigungu_code": _safe_str(cells.get("sigungu_code")),
            "dong_code": _safe_str(cells.get("dong_code")),
            "workplace_type": _safe_int(cells.get("workplace_type")),
            "industry_code": _safe_str(cells.get("industry_code")),
            "industry_name": _safe_str(cells.get("industry_name")),
            "applied_at": _parse_date(cells.get("applied_at")),
            "re_registered_at": _parse_date(cells.get("re_registered_at")),
            "withdrawn_at": _parse_date(cells.get("withdrawn_at")),
            "subscribers": subs,
            "notified_amount": _safe_int(cells.get("notified_amount")),
            "new_acquisitions": _safe_int(cells.get("new_acquisitions")),
            "lost_subscribers": _safe_int(cells.get("lost_subscribers")),
        })

    print(f"파싱 완료: {len(records)}건 대상, {skipped}건 제외 (가입자수 < {min_subscribers} 또는 파싱 실패)")

    if not records:
        print("적재할 데이터가 없습니다.")
        return

    # DB 적재
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()

    try:
        # 1) industries UPSERT (코드만 있고 이름 없는 경우도 포함)
        print("업종 코드 적재 중...")
        industry_map: dict[str, str] = {}
        for r in records:
            code = r["industry_code"]
            if not code:
                continue
            name = r["industry_name"] or f"미분류({code})"
            if code not in industry_map or (r["industry_name"] and industry_map[code].startswith("미분류")):
                industry_map[code] = name

        if industry_map:
            ind_rows = [{"code": c, "name": n} for c, n in industry_map.items()]
            stmt = pg_insert(Industry).values(ind_rows)
            stmt = stmt.on_conflict_do_update(
                index_elements=["code"],
                set_={"name": stmt.excluded.name},
            )
            db.execute(stmt)
        db.flush()
        print(f"  업종: {len(industry_map)}건")

        # 2) workplaces UPSERT (bulk)
        # (biz_no, name) 중복 시 마지막 값만 유지 → 배치 내 ON CONFLICT 충돌 방지
        print("사업장 마스터 적재 중...")
        NPS_RAW_COLS = [
            "enrollment_status", "postal_code", "address_lot", "address_road",
            "legal_dong_code", "admin_dong_code", "sido_code", "sigungu_code",
            "dong_code", "workplace_type", "industry_code",
            "applied_at", "re_registered_at", "withdrawn_at",
        ]

        # in-memory dedup: (biz_no, name) → 최신 row
        wp_dedup: dict[tuple[str, str], dict] = {}
        for r in records:
            if not r["biz_no"]:
                continue
            wp_dedup[(r["biz_no"], r["name"])] = r

        wp_rows = [
            {
                "biz_no": r["biz_no"],
                "name": r["name"],
                **{col: r[col] for col in NPS_RAW_COLS},
            }
            for r in wp_dedup.values()
        ]
        print(f"  사업장 dedup: {len(records)} records → {len(wp_rows)} unique")

        BATCH = 2000  # 16 컬럼 × 2000 rows = 32k params (PG 한도 65535 이내)
        for i in range(0, len(wp_rows), BATCH):
            batch = wp_rows[i : i + BATCH]
            stmt = pg_insert(NpsRaw).values(batch)
            stmt = stmt.on_conflict_do_update(
                constraint="uq_nps_raw_biz_no_name",
                set_={col: getattr(stmt.excluded, col) for col in NPS_RAW_COLS},
            )
            db.execute(stmt)
            db.flush()
            if (i // BATCH) % 10 == 0:
                print(f"  사업장: {min(i + BATCH, len(wp_rows))}/{len(wp_rows)}")
        wp_count = len(wp_rows)

        # (biz_no, name) → id 매핑 조회
        all_wp = db.query(NpsRaw.biz_no, NpsRaw.name, NpsRaw.id).all()
        biz_name_to_id = {(biz_no, name): wid for biz_no, name, wid in all_wp}
        print(f"  사업장 마스터: {wp_count}건 UPSERT")

        # 3) nps_monthly INSERT (bulk, ON CONFLICT DO NOTHING = 이미 있으면 skip)
        print("월별 데이터 적재 중...")
        wm_rows = []
        for r in records:
            wp_id = biz_name_to_id.get((r["biz_no"], r["name"]))
            if wp_id is None:
                continue
            wm_rows.append({
                "nps_id": wp_id,
                "year_month": r["year_month"],
                "subscribers": r["subscribers"],
                "notified_amount": r["notified_amount"],
                "new_acquisitions": r["new_acquisitions"],
                "lost_subscribers": r["lost_subscribers"],
            })

        # 배치 내 (nps_id, year_month) 중복 제거
        wm_dedup: dict[tuple, dict] = {}
        for row in wm_rows:
            wm_dedup[(row["nps_id"], row["year_month"])] = row
        wm_rows = list(wm_dedup.values())
        print(f"  월별 dedup: {len(wm_rows)} rows")

        for i in range(0, len(wm_rows), BATCH):
            batch = wm_rows[i : i + BATCH]
            stmt = pg_insert(NpsMonthly).values(batch)
            stmt = stmt.on_conflict_do_nothing(constraint="uq_nps_monthly_yearmonth")
            db.execute(stmt)
            db.flush()
            if (i // BATCH) % 10 == 0:
                print(f"  월별: {min(i + BATCH, len(wm_rows))}/{len(wm_rows)}")
        wm_count = len(wm_rows)

        db.commit()
        print(f"\n완료! 사업장 {wp_count}건, 월별 {wm_count}건 적재됨.")
        print("서비스 테이블 갱신은: python -m app.compute")

    except Exception as e:
        db.rollback()
        print(f"오류 발생: {e}")
        raise
    finally:
        db.close()


def main():
    parser = argparse.ArgumentParser(description="국민연금 사업장 데이터 → DB 적재")
    parser.add_argument("--file", "-f", required=True, help="CSV 또는 Excel 파일 경로")
    parser.add_argument("--min-subscribers", type=int, default=5, help="최소 가입자수 필터 (기본: 5)")
    args = parser.parse_args()
    ingest(args.file, args.min_subscribers)


if __name__ == "__main__":
    main()
