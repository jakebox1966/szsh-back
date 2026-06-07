"""
엑셀 기업명 → DART corp_code 해석 (비상장 타깃 명단 생성).

corpCode.xml(전체 법인) 에 대해 한영음차·법인키워드·괄호주석·공백·&·숫자·접미사(prefix)·fuzzy 로
강건하게 매칭한 뒤, 비상장(stock_code 없음) 만 추려 data/dart/unlisted_targets.csv 로 저장.
(상장·카테고리헤더는 제외 → 비상장 재무 수집 타깃만 남김)

사용:
  python -m app.resolve_unlisted --xlsx ~/Downloads/top_1000_unlisted_comprehensive.xlsx
"""
import argparse
import csv
import difflib
import re
import xml.etree.ElementTree as ET
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from app.match import _normalize, _hanyoung_canonical

CORPCODE = Path("data/dart/CORPCODE.xml")
OUT = Path("data/dart/unlisted_targets.csv")

# 브랜드성 숫자 음차 (회사명에 흔한 것만)
_NUM = {"11": "십일", "24": "이십사", "365": "삼육오"}


def _loose(name: str) -> str:
    """강건 매칭 키: 괄호주석 제거 + & → 엔 + 숫자음차 + 한영음차 + 공백 제거."""
    n = re.sub(r"\s*\(.*?\)\s*", " ", name or "")  # "SK쉴더스 (ADT캡스)" → "SK쉴더스"
    n = n.replace("&", "엔")                        # E&S → 이엔에스
    for k, v in _NUM.items():
        n = re.sub(rf"(?<!\d){k}(?!\d)", v, n)
    return _hanyoung_canonical(_normalize(n)).replace(" ", "")


def _load_corpcodes():
    root = ET.fromstring(CORPCODE.read_bytes())
    rows = []
    for e in root.iter("list"):
        cc = (e.findtext("corp_code") or "").strip()
        nm = (e.findtext("corp_name") or "").strip()
        sc = (e.findtext("stock_code") or "").strip()
        if nm:
            rows.append((cc, nm, sc))
    return rows


def _excel_names(xlsx: Path):
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    names = set()
    for sn in wb.sheetnames:
        for r in wb[sn].iter_rows(min_row=2, values_only=True):
            if r and r[1]:
                names.add(str(r[1]).strip())
    return sorted(names)


def resolve(xlsx: str):
    corpcodes = _load_corpcodes()
    idx: dict[str, list] = defaultdict(list)
    for cc, nm, sc in corpcodes:
        idx[_loose(nm)].append((cc, nm, sc))
    keys = list(idx.keys())

    names = _excel_names(Path(xlsx).expanduser())

    # 괄호 안이 annotation(자회사·계열·전환 등)이 아니라 실제 법인명인 경우 그것도 후보로
    _ANNO = ("자회사", "계열", "비상장", "사업부", "전환", "상장", "외", "주요", "포함", "기준", "추진")

    def _alt_keys(name):
        keys2 = [_loose(name)]
        for inner in re.findall(r"\(([^)]+)\)", name or ""):
            if not any(a in inner for a in _ANNO):
                keys2.append(_loose(inner))  # "다이소 (아성다이소)" → "아성다이소"
        return [k for k in keys2 if len(k) >= 2]

    def best_candidate(name):
        alts = _alt_keys(name)
        if not alts:
            return None, "too-short"
        cands, how = None, "exact"
        for k in alts:        # 이름·괄호내용 등 후보키로 정확매칭 시도
            if k in idx:
                cands = idx[k]
                break
        k = alts[0]           # prefix/fuzzy 는 주 키(이름) 기준
        if not cands:
            # 접미사 변형: 엑셀키가 corp키의 prefix (AIA생명 → 에이아이에이생명보험)
            pref = [kk for kk in keys if len(k) >= 4 and kk.startswith(k)]
            if pref:
                cands = idx[min(pref, key=len)]
                how = "prefix"
            else:
                # fuzzy (마지막 수단)
                close = difflib.get_close_matches(k, keys, n=1, cutoff=0.9)
                if close:
                    cands = idx[close[0]]
                    how = "fuzzy"
        if not cands:
            return None, "no-match"
        unlisted = sorted([c for c in cands if not c[2]], key=lambda x: len(x[1]))
        if unlisted:
            return unlisted[0], how
        return (None, "listed:" + sorted(cands, key=lambda x: len(x[1]))[0][1])

    targets, listed, failed = {}, [], []
    for name in names:
        pick, how = best_candidate(name)
        if pick:
            targets[pick[0]] = (pick[0], pick[1], name)  # corp_code 기준 dedupe
        elif how.startswith("listed:"):
            listed.append((name, how[7:]))
        else:
            failed.append(name)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["corp_code", "corp_name", "excel_name"])
        for cc, nm, en in targets.values():
            w.writerow([cc, nm, en])

    print(f"엑셀 유니크 기업명: {len(names)}")
    print(f"  → 비상장 타깃(corp 유니크): {len(targets)}곳  저장: {OUT}")
    print(f"  → 상장/헤더(스킵): {len(listed)}곳")
    print(f"  → 해석 실패: {len(failed)}곳")
    print(f"=== 실패 샘플 20 ===\n  {failed[:20]}")
    return targets, listed, failed


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default="~/Downloads/top_1000_unlisted_comprehensive.xlsx")
    resolve(ap.parse_args().xlsx)


if __name__ == "__main__":
    main()
